import sys
sys.path.insert(0, '/tmp/sample-slides-deps')
from pathlib import Path
import cairosvg
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

ROOT = Path('/workspaces/my-codex-codespace')
OUT = ROOT / 'sample-slides'
cairosvg.svg2png(url=str(ROOT/'pelican-on-a-bike.svg'), write_to=str(OUT/'pelican.png'), output_width=1800, output_height=1400)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BG='F4F5E9'; INK='233947'; TEAL='269E9C'; ORANGE='F89B35'; MUTED='607A7F'
def text(s,x,y,w,h,value,size=24,color=INK,bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
    return box

def shape(s,x,y,w,h,color):
    a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=RGBColor.from_string(color); a.line.fill.background()
    return a

def slide(k,title=None,subtitle=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
    text(s,.65,.25,10,.35,'PELICAN CYCLE CLUB  /  SAMPLE PRESENTATION',11,TEAL if not dark else 'FFDA78',True)
    text(s,12,.25,.6,.35,f'{k:02}',11,MUTED if not dark else BG)
    if title: text(s,.65,.9,12,1,title,36,BG if dark else INK,True)
    if subtitle: text(s,.65,1.85,12,.7,subtitle,18,MUTED if not dark else 'DCE7DF')
    text(s,.65,7.05,12,.25,'Illustrative concept • All data and plans are fictional',10,MUTED if not dark else 'DCE7DF')
    return s

s=slide(1)
text(s,.7,1.6,5.6,2.1,'Big beak.\nBigger adventures.',44,INK,True)
text(s,.75,4.1,5.1,1.2,'A playful sample deck for the\nPelican Cycle Club',24,MUTED)
shape(s,.75,5.65,3.4,.6,TEAL)
text(s,.95,5.7,3,.4,'LET’S ROLL',16,'FFFFFF',True)
s.shapes.add_picture(str(OUT/'pelican.png'), Inches(6.1), Inches(1.4), width=Inches(6.8))
s.notes_slide.notes_text_frame.text='Sample title layout. Replace the name, subtitle, and artwork with your own.'

s=slide(2,'A simple idea, with room to grow','Example layout: three pillars')
for x,num,title,body in [( .7,'01','Ride together','Easy routes. Friendly faces.\nA pace everyone can enjoy.'),(4.85,'02','Explore more','Discover coastal paths and\nnew neighborhood favorites.'),(9,'03','Make it fun','A memorable mascot and\na reason to come back.')]:
    shape(s,x,2.8,3.65,3.35,'FFFFFF')
    text(s,x+.25,3.05,3,.5,num,26,TEAL,True)
    text(s,x+.25,3.85,3.1,.6,title,25,INK,True)
    text(s,x+.25,4.65,3.1,1.1,body,19,MUTED)

s=slide(3,'Give your story a memorable face','Example layout: image + supporting message')
s.shapes.add_picture(str(OUT/'pelican.png'),Inches(.7),Inches(2.55),width=Inches(5.65))
text(s,7,2.9,5.3,.6,'Meet your ride companion',28,INK,True)
text(s,7,3.75,5.1,2.3,'A distinctive illustration makes the idea easy to recognize.\n\nUse it on event invitations, welcome slides, and club materials.',22,MUTED)

s=slide(4,'Show momentum at a glance','Example layout: editable chart + key takeaway')
data=CategoryChartData(); data.categories=['April','May','June','July']; data.add_series('Sample riders',[24,38,55,72])
chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.7),Inches(2.8),Inches(7.6),Inches(3.65),data).chart
chart.has_legend=False; chart.has_title=False
chart.chart_style=10
series=chart.series[0]; series.format.fill.solid(); series.format.fill.fore_color.rgb=RGBColor.from_string(TEAL); series.format.line.fill.background()
chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=80
chart.category_axis.tick_labels.font.size=Pt(14); chart.value_axis.tick_labels.font.size=Pt(12)
shape(s,9,2.9,3.6,3.3,'FFFFFF')
text(s,9.3,3.15,3,.9,'3×',58,TEAL,True)
text(s,9.3,4.3,3,1,'Growth in monthly riders',23,INK,True)
text(s,9.3,5.55,3,.4,'Illustrative data only',13,MUTED)

s=slide(5,'Turn the idea into a plan','Example layout: a three-step roadmap')
for x,when,title,body in [( .7,'WEEK 1','Prepare','Choose a route\nCreate an invitation\nConfirm the essentials'),(4.85,'WEEK 2','Launch','Welcome the riders\nRun the first event\nCapture feedback'),(9,'WEEK 3','Improve','Refine the route\nShare the highlights\nPlan the next ride')]:
    shape(s,x,3,3.65,.08,TEAL)
    text(s,x,2.6,3.5,.4,when,14,TEAL,True)
    text(s,x,3.4,3.5,.7,title,30,INK,True)
    text(s,x,4.35,3.5,1.7,body,22,MUTED)

s=slide(6,dark=True)
text(s,.8,1.7,11.8,1.6,'Ready for the next ride?',48,BG,True)
text(s,.85,3.4,10,1.3,'Make the invitation clear.\nMake the experience welcoming.',28,'DCE7DF')
shape(s,.85,5.45,4.1,.75,ORANGE)
text(s,1.05,5.55,3.7,.5,'YOUR CALL TO ACTION',18,INK,True)
prs.save(OUT/'sample-powerpoint.pptx')
# Reopen the generated package to verify slide count and editable chart.
check=Presentation(OUT/'sample-powerpoint.pptx')
assert len(check.slides)==6
assert sum(1 for s in check.slides for a in s.shapes if a.has_chart)==1
print(f'Created and verified {len(check.slides)} slides: {OUT / "sample-powerpoint.pptx"}')
